#!/usr/bin/env python3
"""Count source-text words in every .xliff / .xlf / .sdlxliff / .mqxliff file in a folder.

Usage:
    python xliff_wordcount.py <folder> [-o word_counts.xlsx] [--target] [--recursive]

Output: an Excel file with S.No, File Name, Word Count (plus Segments and Notes),
and a TOTAL row. Works with XLIFF 1.2 (<trans-unit><source>) and XLIFF 2.x
(<unit><segment><source>). Inline code tags (<ph>, <bpt>, <ept>, <it>, <x/>)
are ignored, so only translatable text is counted. CJK characters count
as one word each.
"""
import argparse
import re
import sys
from pathlib import Path
import xml.etree.ElementTree as ET

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

EXTS = {".xliff", ".xlf", ".sdlxliff", ".mqxliff"}
SKIP_TAGS = {"ph", "bpt", "ept", "it", "sc", "ec", "sm", "em"}  # inline codes, not text
CJK = r"\u3040-\u30ff\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff\uac00-\ud7af"
WORD_RE = re.compile(rf"[{CJK}]|[^\s{CJK}]*[\w][^\s{CJK}]*", re.UNICODE)


def local(tag):
    return tag.rsplit("}", 1)[-1] if isinstance(tag, str) else ""


def text_of(elem):
    """Text of an element, skipping inline code tags but keeping their tails."""
    parts = [elem.text or ""]
    for child in elem:
        if local(child.tag) not in SKIP_TAGS:
            parts.append(text_of(child))
        parts.append(child.tail or "")
    return "".join(parts)


def count_words(text):
    return len(WORD_RE.findall(text))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("folder")
    ap.add_argument("-o", "--output", default="word_counts.xlsx")
    ap.add_argument("--target", action="store_true", help="count <target> text instead of <source>")
    ap.add_argument("--recursive", action="store_true", help="include sub-folders")
    a = ap.parse_args()

    folder = Path(a.folder)
    it = folder.rglob("*") if a.recursive else folder.glob("*")
    files = sorted((p for p in it if p.is_file() and p.suffix.lower() in EXTS), key=lambda p: p.name.lower())
    if not files:
        sys.exit(f"No XLIFF files found in {folder}")

    # Count <source> (or <target>) text, ignoring TM matches inside <alt-trans>
    def count_no_alt(p):
        try:
            root = ET.parse(p).getroot()
        except ET.ParseError as e:
            return None, None, f"XML error: {e}"
        want = "target" if a.target else "source"
        alt_children = set()
        for el in root.iter():
            if local(el.tag) == "alt-trans":
                alt_children.update(id(x) for x in el.iter())
        w = s = 0
        for el in root.iter():
            if local(el.tag) == want and id(el) not in alt_children:
                s += 1
                w += count_words(text_of(el))
        return w, s, ""

    wb = Workbook()
    ws = wb.active
    ws.title = "Word Counts"
    headers = ["S.No", "File Name", "Word Count", "Segments", "Notes"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
        c.alignment = Alignment(horizontal="center")

    total = 0
    for i, p in enumerate(files, 1):
        w, s, note = count_no_alt(p)
        name = str(p.relative_to(folder)) if a.recursive else p.name
        ws.append([i, name, w, s, note])
        total += w or 0
        print(f"{i:>4}  {name:<50} {w if w is not None else 'ERROR':>8}")

    n = len(files) + 1
    ws.append([])
    ws.append(["", "TOTAL", f"=SUM(C2:C{n})", f"=SUM(D2:D{n})", ""])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for col, width in zip("ABCDE", (7, 55, 13, 11, 30)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    wb.save(a.output)
    print(f"\n{len(files)} files, {total} words -> {a.output}")


if __name__ == "__main__":
    main()
